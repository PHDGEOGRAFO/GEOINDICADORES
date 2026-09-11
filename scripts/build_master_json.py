#!/usr/bin/env python3
"""Genera JSON maestros del visor desde la BBDD Maestra.

Reglas:
- 0 es dato válido y nunca se convierte en null.
- Solo None o cadenas vacías representan ausencia real.
- La posición de número/nombre de barrio y bloques agregados se detecta,
  evitando depender de columnas fijas A/B o B/C.
- Se exigen exactamente 27 barrios por año.
"""
from __future__ import annotations

import argparse
import json
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

YEARS = (2025, 2026)
EXPECTED_BARRIOS = 27
CODES = ("IAT_", "IUT_", "IST_", "IET_", "IIT_")


def blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def scalar(value: Any) -> Any:
    if blank(value):
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value
    text = str(value).strip()
    try:
        number = float(text.replace(",", "."))
        return int(number) if number.is_integer() else number
    except ValueError:
        return text


def norm(value: Any) -> str:
    return str(value or "").strip().upper()


def find_cell(ws, predicate):
    for row in ws.iter_rows():
        for cell in row:
            if predicate(cell.value):
                return cell
    return None


def code_columns(ws):
    anchor = find_cell(ws, lambda v: norm(v) == "IAT_1")
    if not anchor:
        raise RuntimeError(f"No se encontró IAT_1 en {ws.title}")
    result = {}
    for cell in ws[anchor.row]:
        value = norm(cell.value)
        if value.startswith(CODES):
            result[cell.column] = value
    return result, anchor.row


def detect_barrio_columns(ws, code_row):
    """Detecta columna secuencial 1..27 y la columna de nombre contigua."""
    for seq_col in range(1, min(ws.max_column or 40, 10) + 1):
        values = [ws.cell(code_row + i, seq_col).value for i in range(1, 6)]
        numeric = []
        for value in values:
            try:
                numeric.append(int(float(value)))
            except (TypeError, ValueError):
                numeric.append(None)
        if numeric == [1, 2, 3, 4, 5]:
            name_col = seq_col + 1
            if all(not blank(ws.cell(code_row + i, name_col).value) for i in range(1, 6)):
                return seq_col, name_col
    raise RuntimeError(f"No se pudieron detectar columnas de barrio en {ws.title}")


def parse_barrios(ws):
    codes, code_row = code_columns(ws)
    seq_col, name_col = detect_barrio_columns(ws, code_row)
    records = []
    for r in range(code_row + 1, code_row + EXPECTED_BARRIOS + 1):
        seq = ws.cell(r, seq_col).value
        name = ws.cell(r, name_col).value
        try:
            seq_num = int(float(seq))
        except (TypeError, ValueError):
            raise RuntimeError(f"Hoja {ws.title}: secuencia barrial inválida en fila {r}: {seq!r}")
        if seq_num != len(records) + 1 or blank(name):
            raise RuntimeError(f"Hoja {ws.title}: estructura barrial inesperada en fila {r}")
        vals = {code: scalar(ws.cell(r, col).value) for col, code in codes.items()}
        records.append({"id": seq_num, "nombre": str(name).strip(), "valores": vals})
    return records, list(codes.values())


def parse_block(ws, accepted_types):
    codes, _ = code_columns(ws)
    out = {key: {} for key in accepted_types.values()}
    for row in ws.iter_rows():
        for cell in row:
            tipo = norm(cell.value)
            if tipo not in accepted_types:
                continue
            name = ws.cell(cell.row, cell.column + 1).value
            if blank(name):
                continue
            key = accepted_types[tipo]
            out[key][str(name).strip()] = {
                code: scalar(ws.cell(cell.row, col).value) for col, code in codes.items()
            }
            break
    return out


def parse_territorios(ws):
    return parse_block(ws, {
        "TERRITORIO REAL": "real",
        "TERRITORIO PROMEDIO": "promedio",
        "DIFERENCIA TERRITORIAL": "diferencia",
    })


def parse_comuna(ws):
    return parse_block(ws, {
        "COMUNA REAL": "real",
        "COMUNA PROMEDIO BARRIOS": "promedio_barrios",
    })


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("xlsx", type=Path)
    parser.add_argument("--out-dir", type=Path, default=Path("data"))
    args = parser.parse_args()

    wb = load_workbook(args.xlsx, data_only=True, read_only=False)
    stamp = datetime.now(timezone.utc).isoformat(timespec="seconds")

    master_barrio = {
        "meta": {
            "fuente": args.xlsx.name,
            "generado_utc": stamp,
            "regla_nulos": "0 es dato valido; solo ausencia real se guarda como null",
            "barrios_esperados": EXPECTED_BARRIOS,
        },
        "anios": {},
    }
    master_aggregates = {
        "meta": {
            "fuente": args.xlsx.name,
            "generado_utc": stamp,
            "barrios_esperados": EXPECTED_BARRIOS,
        },
        "anios": {},
    }

    indicator_union = []
    for year in YEARS:
        if str(year) not in wb.sheetnames:
            raise RuntimeError(f"Falta hoja {year}")
        ws = wb[str(year)]
        barrios, codes = parse_barrios(ws)
        if len(barrios) != EXPECTED_BARRIOS:
            raise RuntimeError(f"Hoja {year}: esperados {EXPECTED_BARRIOS}; leídos {len(barrios)}")
        names = [b["nombre"] for b in barrios]
        if len(set(map(norm, names))) != EXPECTED_BARRIOS:
            raise RuntimeError(f"Hoja {year}: nombres de barrio duplicados")
        indicator_union.extend(c for c in codes if c not in indicator_union)
        master_barrio["anios"][str(year)] = barrios
        master_aggregates["anios"][str(year)] = {
            "territorio": parse_territorios(ws),
            "comuna": parse_comuna(ws),
        }

    master_barrio["indicadores"] = indicator_union
    args.out_dir.mkdir(parents=True, exist_ok=True)
    barrio_path = args.out_dir / "master-barrio.json"
    aggregates_path = args.out_dir / "master-aggregates.json"
    barrio_path.write_text(json.dumps(master_barrio, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    aggregates_path.write_text(json.dumps(master_aggregates, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"OK: {barrio_path}")
    print(f"OK: {aggregates_path}")


if __name__ == "__main__":
    main()
